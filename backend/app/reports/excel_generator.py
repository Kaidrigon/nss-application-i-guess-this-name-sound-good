from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen
import os

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    Border,
    Side,
    PatternFill,
)
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as ExcelImage


# ---------------------------------------------------------
# CONFIGURATION
# ---------------------------------------------------------

# Copy your logo into:
# backend/app/reports/assets/nss-logo.png
#
# Or configure NSS_LOGO_PATH to point to a logo file
# accessible by the backend deployment.

DEFAULT_LOGO_PATH = (
    Path(__file__).resolve().parent / "assets" / "nss-logo.png"
)

MAX_EMBEDDED_PHOTOS = 30
IMAGE_TIMEOUT_SECONDS = 8
MAX_IMAGE_BYTES = 5 * 1024 * 1024


# ---------------------------------------------------------
# HELPERS
# ---------------------------------------------------------

def safe_text(value):
    """Convert optional values into readable spreadsheet text."""
    if value is None or value == "":
        return "Not recorded"
    if isinstance(value, list):
        return "\n".join(str(item) for item in value) or "Not recorded"
    return str(value)


def add_section(sheet, row, title, end_column=4):
    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=end_column,
    )

    cell = sheet.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=12, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="176B52")
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24


def style_table_header(sheet, row, headers):
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="247A60")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = THIN_BORDER

    sheet.row_dimensions[row].height = 30


def style_data_cell(cell):
    cell.border = THIN_BORDER
    cell.alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )


def download_image(url):
    """
    Download an HTTPS image for embedding.
    Failed downloads return None so the workbook can still be created.
    """
    if not isinstance(url, str) or not url.startswith("https://"):
        return None

    try:
        request = Request(
            url,
            headers={"User-Agent": "NSS-Report-Generator/1.0"},
        )

        with urlopen(request, timeout=IMAGE_TIMEOUT_SECONDS) as response:
            content_type = response.headers.get("Content-Type", "")
            content_length = response.headers.get("Content-Length")

            if not content_type.lower().startswith("image/"):
                return None

            if content_length and int(content_length) > MAX_IMAGE_BYTES:
                return None

            data = response.read(MAX_IMAGE_BYTES + 1)

            if len(data) > MAX_IMAGE_BYTES:
                return None

            return BytesIO(data)

    except Exception:
        return None


def add_local_logo(sheet):
    """Insert the NSS logo if the configured local file exists."""
    logo_path = Path(
        os.getenv("NSS_LOGO_PATH", str(DEFAULT_LOGO_PATH))
    )

    if not logo_path.is_file():
        return False

    try:
        logo = ExcelImage(str(logo_path))
        logo.width = 75
        logo.height = 75
        sheet.add_image(logo, "A1")
        return True
    except Exception:
        return False


# Shared styles
THIN_SIDE = Side(style="thin", color="D4DED9")
THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)


# ---------------------------------------------------------
# MAIN EXCEL GENERATOR
# ---------------------------------------------------------

def generate_event_excel(report_data: dict) -> BytesIO:
    workbook = Workbook()

    title_font = Font(
        name="Aptos Display",
        bold=True,
        size=18,
        color="176B52",
    )

    subtitle_font = Font(
        italic=True,
        size=10,
        color="555555",
    )

    label_font = Font(bold=True, color="174B3B")

    # =====================================================
    # SHEET 1 — EVENT SUMMARY
    # =====================================================

    sheet = workbook.active
    sheet.title = "Event Summary"
    sheet.sheet_view.showGridLines = False

    # Reserve the first rows for branding.
    sheet.merge_cells("B1:F2")
    sheet["B1"] = "NATIONAL SERVICE SCHEME"
    sheet["B1"].font = title_font
    sheet["B1"].alignment = Alignment(
        vertical="center",
        horizontal="left",
    )

    sheet.merge_cells("B3:F3")
    sheet["B3"] = "NSS EVENT REPORT"
    sheet["B3"].font = Font(bold=True, size=14)
    sheet["B3"].alignment = Alignment(vertical="center")

    sheet.merge_cells("B4:F4")
    sheet["B4"] = "Event Documentation and Participation Record"
    sheet["B4"].font = subtitle_font

    add_local_logo(sheet)

    row = 6
    add_section(sheet, row, "EVENT DETAILS", 6)
    row += 1

    summary_data = [
        ("Event Title", report_data.get("event_title")),
        ("Event Type", report_data.get("event_type")),
        ("Date", report_data.get("date")),
        ("Start Time", report_data.get("start_time")),
        ("End Time", report_data.get("end_time")),
        ("Venue", report_data.get("venue")),
        ("Status", report_data.get("status")),
        ("Description", report_data.get("description")),
        ("Introduction", report_data.get("introduction")),
        ("Event Summary", report_data.get("event_summary")),
        (
            "Faculty In Charge",
            report_data.get("faculty_in_charge"),
        ),
        (
            "Organizing Team",
            report_data.get("organizing_team"),
        ),
        (
            "Objectives",
            report_data.get("objectives"),
        ),
        (
            "Outcomes",
            report_data.get("outcomes"),
        ),
    ]

    for label, value in summary_data:
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = label_font
        label_cell.fill = PatternFill("solid", fgColor="EAF3EE")
        style_data_cell(label_cell)

        sheet.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=6,
        )

        value_cell = sheet.cell(
            row=row,
            column=2,
            value=safe_text(value),
        )
        style_data_cell(value_cell)

        # Give longer text enough room.
        text_length = len(str(value or ""))
        sheet.row_dimensions[row].height = min(
            90,
            max(24, 18 * (text_length // 75 + 1)),
        )

        row += 1

    row += 1
    add_section(sheet, row, "PARTICIPATION & SERVICE HOURS", 6)
    row += 1

    participation = [
        ("Registered Volunteers", report_data.get("registered_volunteers")),
        ("Attended Volunteers", report_data.get("attended_volunteers")),
        ("Absent Volunteers", report_data.get("absent_volunteers")),
        ("Attendance Percentage", report_data.get("attendance_percentage")),
        (
            "Service Hours per Volunteer",
            report_data.get("credited_hours_per_volunteer"),
        ),
        ("Total Service Hours", report_data.get("total_service_hours")),
        ("Male Attendees", report_data.get("male_attendees")),
        ("Female Attendees", report_data.get("female_attendees")),
        (
            "Other Recorded Gender Attendees",
            report_data.get("other_gender_attendees"),
        ),
        (
            "Attendees with Unrecorded Gender",
            report_data.get("unrecorded_gender_attendees"),
        ),
        ("Faculty Participants", report_data.get("faculty_participants")),
        ("Beneficiaries", report_data.get("beneficiaries")),
    ]

    for label, value in participation:
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = label_font
        label_cell.fill = PatternFill("solid", fgColor="EAF3EE")
        style_data_cell(label_cell)

        sheet.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=3,
        )

        display_value = (
            "Not recorded" if value is None else value
        )

        value_cell = sheet.cell(
            row=row,
            column=2,
            value=display_value,
        )
        style_data_cell(value_cell)

        row += 1

    # Format attendance percentage.
    for cells in sheet.iter_rows():
        for cell in cells:
            if cell.value == "Attendance Percentage":
                percentage_cell = sheet.cell(
                    row=cell.row,
                    column=2,
                )
                percentage_cell.number_format = '0.00"%"'

    for column, width in {
        "A": 32,
        "B": 24,
        "C": 20,
        "D": 20,
        "E": 20,
        "F": 20,
    }.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A6"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"

    # =====================================================
    # SHEET 2 — YEAR BREAKDOWN
    # =====================================================

    year_sheet = workbook.create_sheet("Year Breakdown")
    year_sheet.sheet_view.showGridLines = False

    year_sheet.merge_cells("A1:D1")
    year_sheet["A1"] = "YEAR-WISE VOLUNTEER PARTICIPATION"
    year_sheet["A1"].font = title_font

    year_headers = [
        "Year",
        "Registered Volunteers",
        "Attended Volunteers",
        "Absent Volunteers",
    ]
    style_table_header(year_sheet, 3, year_headers)

    row = 4

    for item in report_data.get("year_breakdown", []):
        values = [
            item.get("year"),
            item.get("registered_volunteers", 0),
            item.get("attended_volunteers", 0),
            item.get("absent_volunteers", 0),
        ]

        for column, value in enumerate(values, start=1):
            cell = year_sheet.cell(
                row=row,
                column=column,
                value=value,
            )
            style_data_cell(cell)

        row += 1

    if row == 4:
        year_sheet["A4"] = "No year-wise records available."

    for column in range(1, 5):
        year_sheet.column_dimensions[
            get_column_letter(column)
        ].width = 26

    year_sheet.freeze_panes = "A4"

    # =====================================================
    # SHEET 3 — CLASS BREAKDOWN
    # =====================================================

    class_sheet = workbook.create_sheet("Class Breakdown")
    class_sheet.sheet_view.showGridLines = False

    class_sheet.merge_cells("A1:D1")
    class_sheet["A1"] = "CLASS-WISE VOLUNTEER PARTICIPATION"
    class_sheet["A1"].font = title_font

    class_headers = [
        "Class",
        "Registered Volunteers",
        "Attended Volunteers",
        "Absent Volunteers",
    ]
    style_table_header(class_sheet, 3, class_headers)

    row = 4

    for item in report_data.get("class_breakdown", []):
        values = [
            item.get("class_name", "Unknown"),
            item.get("registered_volunteers", 0),
            item.get("attended_volunteers", 0),
            item.get("absent_volunteers", 0),
        ]

        for column, value in enumerate(values, start=1):
            cell = class_sheet.cell(
                row=row,
                column=column,
                value=value,
            )
            style_data_cell(cell)

        row += 1

    if row == 4:
        class_sheet["A4"] = "No class-wise records available."

    for column in range(1, 5):
        class_sheet.column_dimensions[
            get_column_letter(column)
        ].width = 28

    class_sheet.freeze_panes = "A4"

    # =====================================================
    # SHEET 4 — EVIDENCE
    # =====================================================

    evidence_sheet = workbook.create_sheet("Evidence")
    evidence_sheet.sheet_view.showGridLines = False

    evidence_sheet.merge_cells("A1:F1")
    evidence_sheet["A1"] = "EVENT EVIDENCE AND PHOTOGRAPHS"
    evidence_sheet["A1"].font = title_font

    evidence_headers = [
        "Photo",
        "File Name",
        "Source URL",
        "ImageKit File ID",
        "Uploaded By",
        "Uploaded At",
    ]
    style_table_header(evidence_sheet, 3, evidence_headers)

    row = 4
    photos = report_data.get("evidence_photos", [])

    for index, photo in enumerate(photos):
        file_url = photo.get("file_url", "")

        # Keep the URL available even when embedding is impossible.
        evidence_sheet.cell(row=row, column=2, value=photo.get("file_name", ""))

        url_cell = evidence_sheet.cell(
            row=row,
            column=3,
            value=file_url or "URL not recorded",
        )

        if file_url.startswith("https://"):
            url_cell.hyperlink = file_url
            url_cell.style = "Hyperlink"

        evidence_sheet.cell(
            row=row,
            column=4,
            value=photo.get("imagekit_file_id", ""),
        )
        evidence_sheet.cell(
            row=row,
            column=5,
            value=photo.get("uploaded_by", ""),
        )
        evidence_sheet.cell(
            row=row,
            column=6,
            value=photo.get("uploaded_at", ""),
        )

        for column in range(2, 7):
            style_data_cell(evidence_sheet.cell(row=row, column=column))

        # Avoid creating an excessively large workbook.
        if index < MAX_EMBEDDED_PHOTOS:
            image_buffer = download_image(file_url)

            if image_buffer:
                try:
                    image = ExcelImage(image_buffer)
                    image.width = 135
                    image.height = 90
                    evidence_sheet.add_image(image, f"A{row}")
                except Exception:
                    evidence_sheet.cell(
                        row=row,
                        column=1,
                        value="Preview unavailable",
                    )
            else:
                evidence_sheet.cell(
                    row=row,
                    column=1,
                    value="Preview unavailable",
                )
        else:
            evidence_sheet.cell(
                row=row,
                column=1,
                value="See source URL",
            )

        evidence_sheet.row_dimensions[row].height = 100
        row += 1

    if not photos:
        evidence_sheet.merge_cells("A4:F4")
        evidence_sheet["A4"] = "No event evidence photos have been recorded."
        evidence_sheet["A4"].alignment = Alignment(wrap_text=True)

    for column, width in {
        "A": 22,
        "B": 30,
        "C": 55,
        "D": 30,
        "E": 25,
        "F": 28,
    }.items():
        evidence_sheet.column_dimensions[column].width = width

    evidence_sheet.freeze_panes = "A4"

    # =====================================================
    # RETURN WORKBOOK
    # =====================================================

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output