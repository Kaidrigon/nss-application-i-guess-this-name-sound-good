
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


# Optional NSS logo.
DEFAULT_LOGO_PATH = (
    Path(__file__).resolve().parent / "assets" / "nss-logo.png"
)

MAX_EMBEDDED_PHOTOS = 30
IMAGE_TIMEOUT_SECONDS = 8
MAX_IMAGE_BYTES = 5 * 1024 * 1024


def display_text(value):
    """Convert report values into readable Word-document text."""
    if value is None or value == "":
        return "Not recorded"

    if isinstance(value, list):
        return "\n".join(str(item) for item in value) or "Not recorded"

    return str(value)


def add_section_heading(document, heading):
    """Add a consistent heading for each report section."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run(heading.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(23, 107, 82)

    return paragraph


def add_key_value_table(document, rows):
    """Create a two-column table for labels and their values."""
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = display_text(value)

        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for run in cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(23, 75, 59)

    document.add_paragraph()
    return table


def add_breakdown_table(document, title, headers, records):
    """Create a table for year-wise or class-wise statistics."""
    add_section_heading(document, title)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header

        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # A simple, readable header background.
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "176B52")
        cell._tc.get_or_add_tcPr().append(shading)

    if not records:
        cells = table.add_row().cells
        cells[0].text = "No records available."
        if len(cells) > 1:
            cells[0].merge(cells[-1])
    else:
        for record in records:
            cells = table.add_row().cells

            for index, value in enumerate(record):
                cells[index].text = display_text(value)

    document.add_paragraph()


def download_report_image(url):
    """
    Download an HTTPS image for the report.
    Returns None if the image cannot be retrieved.
    """
    if not isinstance(url, str) or not url.startswith("https://"):
        return None

    try:
        request = Request(
            url,
            headers={"User-Agent": "NSS-Report-Generator/1.0"},
        )

        with urlopen(request, timeout=IMAGE_TIMEOUT_SECONDS) as response:
            content_type = response.headers.get("Content-Type", "")
            content_length = response.headers.get("Content-Length")

            if not content_type.lower().startswith("image/"):
                return None

            if content_length:
                try:
                    if int(content_length) > MAX_IMAGE_BYTES:
                        return None
                except ValueError:
                    return None

            data = response.read(MAX_IMAGE_BYTES + 1)

            if len(data) > MAX_IMAGE_BYTES:
                return None

            return BytesIO(data)

    except Exception:
        return None


def add_evidence_photos(document, photos):
    """Add available event photos and retain source links if embedding fails."""
    add_section_heading(document, "Event Evidence and Photographs")

    if not photos:
        document.add_paragraph("No event evidence photos have been recorded.")
        return

    for index, photo in enumerate(photos):
        if index >= MAX_EMBEDDED_PHOTOS:
            document.add_paragraph(
                "Additional photos are available in the event gallery."
            )
            break

        filename = photo.get("file_name") or f"Event photo {index + 1}"
        document.add_paragraph(filename, style="Caption")

        image_buffer = download_report_image(photo.get("file_url", ""))

        if image_buffer:
            try:
                document.add_picture(image_buffer, width=Inches(4.8))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                document.add_paragraph(
                    "Image preview unavailable. Use the source link below."
                )
        else:
            document.add_paragraph(
                "Image preview unavailable. Use the source link below."
            )

        file_url = photo.get("file_url", "")
        if file_url.startswith("https://"):
            link_paragraph = document.add_paragraph()
            link_paragraph.add_run("Open original photo: ").bold = True
            link_paragraph.add_run(file_url)

        metadata = []
        if photo.get("uploaded_by"):
            metadata.append(f"Uploaded by: {photo['uploaded_by']}")
        if photo.get("uploaded_at"):
            metadata.append(f"Uploaded at: {photo['uploaded_at']}")

        if metadata:
            document.add_paragraph(" | ".join(metadata))


def generate_event_word(report_data: dict) -> BytesIO:
    """Generate a downloadable Word report from existing report data."""
    document = Document()

    # Default document typography.
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    # Optional logo.
    logo_path = Path(
        os.getenv("NSS_LOGO_PATH", str(DEFAULT_LOGO_PATH))
    )

    if logo_path.is_file():
        try:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(
                str(logo_path), width=Inches(1.0)
            )
        except Exception:
            pass

    # Cover/title area.
    title = document.add_heading("NATIONAL SERVICE SCHEME", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph("NSS EVENT REPORT")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(23, 107, 82)

    event_title = report_data.get("event_title") or "Untitled Event"
    event_heading = document.add_paragraph()
    event_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    event_run = event_heading.add_run(event_title)
    event_run.bold = True
    event_run.font.size = Pt(14)

    document.add_paragraph(
        "Event Documentation and Participation Record"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1: event details and narrative.
    add_section_heading(document, "Event Details")
    add_key_value_table(document, [
        ("Event Title", report_data.get("event_title")),
        ("Event Type", report_data.get("event_type")),
        ("Date", report_data.get("date")),
        ("Start Time", report_data.get("start_time")),
        ("End Time", report_data.get("end_time")),
        ("Venue", report_data.get("venue")),
        ("Status", report_data.get("status")),
        ("Description", report_data.get("description")),
        ("Introduction", report_data.get("introduction")),
        ("Event Summary", report_data.get("event_summary")),
        ("Faculty In Charge", report_data.get("faculty_in_charge")),
        ("Organizing Team", report_data.get("organizing_team")),
        ("Objectives", report_data.get("objectives")),
        ("Outcomes", report_data.get("outcomes")),
    ])

    # Section 2: participation and service hours.
    add_section_heading(document, "Participation and Service Hours")
    add_key_value_table(document, [
        ("Registered Volunteers", report_data.get("registered_volunteers")),
        ("Attended Volunteers", report_data.get("attended_volunteers")),
        ("Absent Volunteers", report_data.get("absent_volunteers")),
        ("Attendance Percentage",
            (
                f"{report_data['attendance_percentage']:.2f}%"
                if report_data.get("attendance_percentage") is not None
                else None
            )),
        ("Service Hours per Volunteer",
         report_data.get("credited_hours_per_volunteer")),
        ("Total Service Hours", report_data.get("total_service_hours")),
        ("Male Attendees", report_data.get("male_attendees")),
        ("Female Attendees", report_data.get("female_attendees")),
        ("Other Recorded Gender Attendees",
         report_data.get("other_gender_attendees")),
        ("Attendees with Unrecorded Gender",
         report_data.get("unrecorded_gender_attendees")),
        ("Faculty Participants", report_data.get("faculty_participants")),
        ("Beneficiaries", report_data.get("beneficiaries")),
    ])

    # Section 3: year-wise breakdown.
    add_breakdown_table(
        document,
        "Year-wise Volunteer Participation",
        ["Year", "Registered", "Attended", "Absent"],
        [
            [
                item.get("year"),
                item.get("registered_volunteers", 0),
                item.get("attended_volunteers", 0),
                item.get("absent_volunteers", 0),
            ]
            for item in report_data.get("year_breakdown", [])
        ],
    )

    # Section 4: class-wise breakdown.
    add_breakdown_table(
        document,
        "Class-wise Volunteer Participation",
        ["Class", "Registered", "Attended", "Absent"],
        [
            [
                item.get("class_name", "Unknown"),
                item.get("registered_volunteers", 0),
                item.get("attended_volunteers", 0),
                item.get("absent_volunteers", 0),
            ]
            for item in report_data.get("class_breakdown", [])
        ],
    )

    # Section 5: photo evidence.
    add_evidence_photos(
        document,
        report_data.get("evidence_photos", []),
    )

    # Return the finished document in memory.
    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output
